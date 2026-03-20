import re
from io import StringIO

import pandas as pd
from lxml import etree

from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR
from pptx.oxml.ns import qn
from pptx.enum.shapes import MSO_SHAPE

from Components.utils import pt, rgb, set_no_fill, set_solid_fill, set_no_line, to_hex
from Components.text_tools import render_html_into_shape


# ── Helpers ────────────────────────────────────────────────────────────────────

def _split_row(row: str) -> list[str]:
    row = row.strip().strip("|")
    parts = re.split(r"(?<!\\)\|", row)
    return [cell.replace("\\|", "|").strip() for cell in parts]


def _parse_markdown_table(md: str) -> list[list[str]]:
    lines = [line for line in md.splitlines() if line.strip()]
    if len(lines) < 2:
        return []
    header = _split_row(lines[0])
    data_lines = lines[2:] if set(lines[1].replace("|", "").strip("-").strip()) == set() else lines[1:]
    rows = [_split_row(line) for line in data_lines]
    return [header] + rows


def _color_from_style(value: str, fallback: RGBColor) -> RGBColor:
    if not isinstance(value, str):
        return fallback
    hex_value = value.strip()
    if hex_value.startswith("#"):
        hex_value = hex_value[1:]
    try:
        if len(hex_value) == 6:
            r = int(hex_value[0:2], 16)
            g = int(hex_value[2:4], 16)
            b = int(hex_value[4:6], 16)
            return RGBColor(r, g, b)
    except ValueError:
        pass
    return fallback


def _set_cell_border(cell, color: RGBColor, width_pt: float = 1.0) -> None:
    """Set all four borders of a table cell via XML (python-pptx has no high-level API)."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    # Insert borders at the BEGINNING of tcPr to ensure they are not shadowed by solidFill
    # Order in tcPr schema: lnL, lnR, lnT, lnB ...
    for i, edge in enumerate(("lnL", "lnR", "lnT", "lnB")):
        # Remove if existing to avoid duplicates
        existing = tcPr.find(qn(f"a:{edge}"))
        if existing is not None:
            tcPr.remove(existing)
        
        ln = etree.Element(qn(f"a:{edge}"))
        tcPr.insert(i, ln) # Insert at the start in order
        ln.set("w", str(int(width_pt * 12700)))
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        
        solidFill = etree.SubElement(ln, qn("a:solidFill"))
        srgbClr = etree.SubElement(solidFill, qn("a:srgbClr"))
        srgbClr.set("val", to_hex(color).upper())
        
        prstDash = etree.SubElement(ln, qn("a:prstDash"))
        prstDash.set("val", "solid")
        
        round_elem = etree.SubElement(ln, qn("a:round"))
        
        headEnd = etree.SubElement(ln, qn("a:headEnd"))
        headEnd.set("type", "none")
        headEnd.set("w", "med")
        headEnd.set("len", "med")
        
        tailEnd = etree.SubElement(ln, qn("a:tailEnd"))
        tailEnd.set("type", "none")
        tailEnd.set("w", "med")
        tailEnd.set("len", "med")


def _lines_from_value(value: object) -> list[str]:
    if isinstance(value, list):
        flattened = "\n".join(str(v) for v in value)
    else:
        flattened = str(value)
    normalized = (
        flattened.replace("<br />", "\n")
        .replace("<br/>", "\n")
        .replace("<br>", "\n")
    )
    return [line for line in normalized.splitlines() if line.strip()] or [""]


# ── Core table renderer ────────────────────────────────────────────────────────

def _render_table_core(
    slide_object,
    component: dict,
    x: float | None,
    y: float | None,
    width: float | None,
    height: float | None,
    header_bg: RGBColor,
    header_text: RGBColor,
    border_color: RGBColor,
    stripe_even: RGBColor | None,
    stripe_odd: RGBColor | None,
    header_bold: bool = True,
    body_bold: bool = False,
    font_size: int = 11,
    cell_bg: list[RGBColor] | None = None,
    cell_text_color: list[RGBColor] | None = None,
) -> None:
    content = component.get("content", "")
    rows: list[list[str]] = []
    if isinstance(content, list):
        rows = content
    elif isinstance(content, str) and content.strip():
        rows = _parse_markdown_table(content)
    if not rows:
        return

    styles = component.get("styles", {}) if isinstance(component.get("styles"), dict) else {}
    font_size = styles.get("fontSize", font_size)

    num_rows = len(rows)
    num_cols = max(len(r) for r in rows)

    if width is None or height is None or x is None or y is None:
        width = slide_object.chart_width
        height = slide_object.get_chart_height()
        x, y = slide_object.get_next_chart_position(height)
    max_height_style = styles.get("maxHeight")
    if isinstance(max_height_style, (int, float)) and max_height_style > 0:
        height = min(height, float(max_height_style))

    custom_widths = component.get("column_widths") or styles.get("ratio")
    if isinstance(custom_widths, list) and custom_widths and all(
        isinstance(v, (int, float)) and v > 0 for v in custom_widths
    ):
        ratios = custom_widths[:num_cols]
        if len(ratios) < num_cols:
            ratios += [1.0] * (num_cols - len(ratios))
        total = sum(ratios) or num_cols
        col_widths = [max(40.0, width * r / total) for r in ratios]
    else:
        col_widths = [width / num_cols] * num_cols

    # Estimate row heights
    estimated_heights: list[float] = []
    for row in rows:
        line_count = max(len(_lines_from_value(cell)) for cell in row) if row else 1
        estimated = max(16.0, line_count * font_size * 1.15 + 4)
        estimated_heights.append(estimated)

    total_estimated = sum(estimated_heights)
    if total_estimated <= height or total_estimated == 0:
        row_heights = estimated_heights
    else:
        scale = height / total_estimated
        row_heights = [max(14.0, h * scale) for h in estimated_heights]

    # python-pptx add_table expects EMU for x, y, width, height; and row/col counts
    slide = slide_object.slide

    # We use the sum of row_heights for the table's height to prevent stretching
    total_h = sum(row_heights)
    table_frame = slide.shapes.add_table(
        num_rows,
        num_cols,
        Pt(x),
        Pt(y),
        width=Pt(width),
        height=Pt(total_h),
    )
    table = table_frame.table
    table.first_row = False
    table.horz_banding = False
    table.vert_banding = False

    # Set column widths
    for c, cw in enumerate(col_widths):
        table.columns[c].width = pt(cw)

    # Set row heights
    for r, rh in enumerate(row_heights):
        table.rows[r].height = pt(rh)

    anchor_value = str(styles.get("anchor", "center") or "").lower()
    autofit_value = styles.get("autofit", "shape")
    align_value = str(styles.get("align", "left") or "").lower()

    for r, row in enumerate(rows):
        for c in range(num_cols):
            cell = table.cell(r, c)
            cell.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            tf.margin_left = Pt(8)
            tf.margin_right = Pt(8)
            tf.margin_top = Pt(5)
            tf.margin_bottom = Pt(5)

            # ── Background ────────────────────────────
            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = header_bg
            else:
                if cell_bg:
                    cell_index = ((r - 1) * num_cols + c) % len(cell_bg)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = cell_bg[cell_index]
                elif stripe_even is None and stripe_odd is None:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(255, 255, 255)
                else:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = stripe_even if r % 2 == 0 else (stripe_odd or RGBColor(255, 255, 255))

            # ── Borders ───────────────────────────────
            _set_cell_border(cell, border_color, 1.0)

            # ── Text content ──────────────────────────
            # clear default paragraphs/runs
            for i in range(len(tf.paragraphs) - 1, 0, -1):
                p = tf.paragraphs[i]._p
                p.getparent().remove(p)
            for run in tf.paragraphs[0].runs:
                run._r.getparent().remove(run._r)

            cell_value = row[c] if c < len(row) else ""
            if isinstance(cell_value, list):
                lines = [str(item) for item in cell_value]
            else:
                text_val = str(cell_value).replace("<br />", "\n").replace("<br/>", "\n").replace("<br>", "\n")
                lines = text_val.splitlines() or [""]

            base_bold = header_bold if r == 0 else body_bold
            base_color = header_text if r == 0 else RGBColor(0, 0, 0)
            if r > 0 and cell_text_color:
                idx = ((r - 1) * num_cols + c) % len(cell_text_color)
                base_color = cell_text_color[idx]

            first_line = True
            for line in lines:
                if first_line:
                    para = tf.paragraphs[0]
                    first_line = False
                else:
                    para = tf.add_paragraph()

                if align_value == "right":
                    para.alignment = PP_ALIGN.RIGHT
                elif align_value == "center" or r == 0: # Center headers by default
                    para.alignment = PP_ALIGN.CENTER
                else:
                    para.alignment = PP_ALIGN.LEFT

                # Bullet for markdown-style items
                stripped = line.lstrip()
                if stripped.startswith("- "):
                    pPr = para._pPr
                    if pPr is None:
                        pPr = para._p.get_or_add_pPr()
                    buChar = etree.SubElement(pPr, qn("a:buChar"))
                    buChar.set("char", "•")
                    line = stripped[2:].lstrip()

                parts = re.split(r"(\*\*.*?\*\*)", line)
                for part in parts:
                    if not part:
                        continue
                    is_bold = part.startswith("**") and part.endswith("**") and len(part) >= 4
                    text_part = part[2:-2] if is_bold else part
                    run = para.add_run()
                    run.text = text_part
                    run.font.size = Pt(font_size)
                    run.font.bold = base_bold or is_bold
                    run.font.color.rgb = base_color

    slide_object.last_bottom_y = max(slide_object.last_bottom_y, y + height)


# ── Public render functions ────────────────────────────────────────────────────

def render_table(
    slide_object,
    component: dict,
    x: float | None = None,
    y: float | None = None,
    width: float | None = None,
    height: float | None = None,
    header_bg: RGBColor | None = None,
    header_text: RGBColor | None = None,
    border_color: RGBColor | None = None,
) -> None:
    """Render a table component from HTML or markdown."""
    content = component.get("content", "")
    if not isinstance(content, str) or not content.strip():
        return

    styles = component.get("styles", {}) if isinstance(component.get("styles"), dict) else {}
    font_size = styles.get("font_size") or styles.get("fontSize") or 11

    if width is None or height is None or x is None or y is None:
        width = slide_object.chart_width
        height = slide_object.get_chart_height()
        x, y = slide_object.get_next_chart_position(height)
    max_height_style = styles.get("maxHeight")
    if isinstance(max_height_style, (int, float)) and max_height_style > 0:
        height = min(height, float(max_height_style))

    rows: list[list[str]] | None = None
    try:
        dfs = pd.read_html(StringIO(content))
        if dfs:
            df = dfs[0]
            rows = [list(df.columns)]
            rows.extend(df.astype(str).fillna("").values.tolist())
    except Exception:
        rows = None

    if rows is None:
        md_rows = _parse_markdown_table(content)
        rows = md_rows if md_rows else None

    if rows:
        header_bg_color = _color_from_style(
            styles.get("header_bg"), header_bg or RGBColor(240, 244, 252)
        )
        header_text_color = _color_from_style(
            styles.get("header_text"), header_text or RGBColor(16, 32, 94)
        )
        border_c = _color_from_style(
            styles.get("border_color"), border_color or RGBColor(200, 200, 200)
        )
        cell_bg_values = styles.get("cell_bg")
        cell_bg_list: list[RGBColor] | None = None
        if isinstance(cell_bg_values, list):
            parsed = [
                _color_from_style(val, RGBColor(255, 255, 255))
                for val in cell_bg_values if isinstance(val, str)
            ]
            if parsed:
                cell_bg_list = parsed
        cell_text_color_values = styles.get("cell_text_color")
        cell_text_color_list: list[RGBColor] | None = None
        if isinstance(cell_text_color_values, list):
            parsed_text = [
                _color_from_style(val, RGBColor(0, 0, 0))
                for val in cell_text_color_values if isinstance(val, str)
            ]
            if parsed_text:
                cell_text_color_list = parsed_text
        stripe_even = RGBColor(250, 251, 253)
        stripe_odd = RGBColor(255, 255, 255)
        _render_table_core(
            slide_object,
            component | {"content": rows},
            x, y, width, height,
            header_bg_color, header_text_color, border_c,
            stripe_even, stripe_odd,
            header_bold=True, body_bold=False,
            font_size=font_size,
            cell_bg=cell_bg_list,
            cell_text_color=cell_text_color_list,
        )
        return

    # Fallback: render raw HTML into a text-box shape
    slide = slide_object.slide
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Pt(x),
        Pt(y),
        Pt(width),
        Pt(height),
    )
    set_no_fill(shape)
    set_no_line(shape)
    render_html_into_shape(shape, content)


def render_meeting_info_table(
    slide_object,
    component: dict,
    x: float | None = None,
    y: float | None = None,
    width: float | None = None,
    height: float | None = None,
) -> None:
    """Render meeting-info table from markdown only (no HTML parsing)."""
    content = component.get("content", "")
    if not isinstance(content, str) or not content.strip():
        return

    if width is None or height is None or x is None or y is None:
        width = slide_object.chart_width
        height = slide_object.get_chart_height()
        x, y = slide_object.get_next_chart_position(height)

    header_bg = RGBColor(33, 45, 106)
    header_text = RGBColor(255, 255, 255)
    border_color = RGBColor(255, 255, 255)
    stripe_even = RGBColor(232, 232, 235)
    stripe_odd = RGBColor(204, 205, 212)

    _render_table_core(
        slide_object,
        component,
        x, y, width, height,
        header_bg, header_text, border_color,
        stripe_even, stripe_odd,
        header_bold=True, body_bold=False,
        font_size=9,
    )
